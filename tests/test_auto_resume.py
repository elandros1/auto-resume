"""Tests for auto_resume package."""

import json

import pytest
from docx import Document

from auto_resume.detector import FieldDetector
from auto_resume.engine import TemplateEngine
from auto_resume.models import Education, ResumeData
from auto_resume.template_generator import generate_all_templates

# ──────────────────── Fixtures ────────────────────

@pytest.fixture
def sample_resume_data():
    """Sample resume data dict."""
    return {
        "name": "测试用户",
        "gender": "男",
        "birth_date": "1995-01",
        "phone": "13900000000",
        "email": "test@test.com",
        "address": "测试地址",
        "hometown": "测试籍贯",
        "political_status": "群众",
        "id_number": "110000000000000000",
        "nationality": "汉",
        "marital_status": "未婚",
        "self_evaluation": "测试自我评价",
        "expected_position": "软件工程师",
        "expected_salary": "20k",
        "expected_city": "北京",
        "availability": "随时",
        "education": [
            {
                "school": "测试大学",
                "major": "计算机科学",
                "degree": "硕士",
                "start_date": "2018.09",
                "end_date": "2021.06",
                "gpa": "3.8",
                "description": "测试描述",
            }
        ],
        "work_experience": [
            {
                "company": "测试公司",
                "position": "开发工程师",
                "start_date": "2021.07",
                "end_date": "至今",
                "department": "技术部",
                "description": "负责后端开发",
            }
        ],
        "projects": [
            {
                "name": "测试项目",
                "role": "负责人",
                "start_date": "2022.01",
                "end_date": "2022.12",
                "description": "项目描述",
                "technologies": "Python",
            }
        ],
        "skills": [
            {"category": "编程", "items": "Python, Java"}
        ],
        "certificates": ["CET-6"],
        "languages": ["英语"],
        "hobbies": ["阅读"],
    }


@pytest.fixture
def sample_resume(sample_resume_data):
    return ResumeData.from_dict(sample_resume_data)


@pytest.fixture
def resume_json_file(tmp_path, sample_resume_data):
    path = tmp_path / "resume.json"
    with open(path, "w", encoding="utf-8") as f:
        json.dump(sample_resume_data, f, ensure_ascii=False)
    return path


@pytest.fixture
def generated_templates(tmp_path):
    """Generate sample templates."""
    templates_dir = tmp_path / "templates"
    generate_all_templates(templates_dir)
    return templates_dir


# ──────────────────── Model Tests ────────────────────

class TestResumeData:
    def test_from_dict_basic(self, sample_resume_data):
        resume = ResumeData.from_dict(sample_resume_data)
        assert resume.name == "测试用户"
        assert resume.gender == "男"
        assert resume.phone == "13900000000"
        assert resume.email == "test@test.com"

    def test_from_dict_education(self, sample_resume_data):
        resume = ResumeData.from_dict(sample_resume_data)
        assert len(resume.education) == 1
        assert resume.education[0].school == "测试大学"
        assert resume.education[0].degree == "硕士"

    def test_from_dict_work(self, sample_resume_data):
        resume = ResumeData.from_dict(sample_resume_data)
        assert len(resume.work_experience) == 1
        assert resume.work_experience[0].company == "测试公司"

    def test_from_json(self, resume_json_file):
        resume = ResumeData.from_json(resume_json_file)
        assert resume.name == "测试用户"

    def test_to_flat_dict_basic(self, sample_resume):
        d = sample_resume.to_flat_dict()
        assert d["name"] == "测试用户"
        assert d["phone"] == "13900000000"
        assert d["email"] == "test@test.com"

    def test_to_flat_dict_education(self, sample_resume):
        d = sample_resume.to_flat_dict()
        assert d["education_1_school"] == "测试大学"
        assert d["education_1_degree"] == "硕士"
        assert d["education_count"] == "1"

    def test_to_flat_dict_work(self, sample_resume):
        d = sample_resume.to_flat_dict()
        assert d["work_1_company"] == "测试公司"
        assert d["work_experience_count"] == "1"

    def test_to_flat_dict_skills(self, sample_resume):
        d = sample_resume.to_flat_dict()
        assert d["skill_1_category"] == "编程"
        assert d["skills_summary"] == "编程: Python, Java"

    def test_to_flat_dict_certificates(self, sample_resume):
        d = sample_resume.to_flat_dict()
        assert d["certificates_summary"] == "CET-6"
        assert d["certificate_1"] == "CET-6"

    def test_empty_resume(self):
        resume = ResumeData()
        d = resume.to_flat_dict()
        assert d["name"] == ""
        assert d["education_count"] == "0"

    def test_education_to_text(self):
        edu = Education(
            school="清华", major="CS", degree="博士",
            start_date="2019", end_date="2024"
        )
        text = edu.to_text()
        assert "清华" in text
        assert "CS" in text
        assert "博士" in text


# ──────────────────── Engine Tests ────────────────────

class TestTemplateEngine:
    def test_fill_placeholder(self, sample_resume, tmp_path):
        """Test filling {{placeholder}} style template."""
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("姓名：{{name}}，电话：{{phone}}")

        template_path = tmp_path / "test_template.docx"
        doc.save(str(template_path))

        engine = TemplateEngine(sample_resume)
        output_path = tmp_path / "output.docx"
        engine.fill_template(template_path, output_path)

        # Read back and verify
        result_doc = Document(str(output_path))
        assert "测试用户" in result_doc.paragraphs[0].text
        assert "13900000000" in result_doc.paragraphs[0].text

    def test_fill_dollar_brace(self, sample_resume, tmp_path):
        """Test filling ${placeholder} style template."""
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("邮箱：${email}")

        template_path = tmp_path / "test_dollar.docx"
        doc.save(str(template_path))

        engine = TemplateEngine(sample_resume)
        output_path = tmp_path / "output_dollar.docx"
        engine.fill_template(template_path, output_path)

        result_doc = Document(str(output_path))
        assert "test@test.com" in result_doc.paragraphs[0].text

    def test_fill_table(self, sample_resume, tmp_path):
        """Test filling a table-based template."""
        doc = Document()
        table = doc.add_table(rows=2, cols=4)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "姓名"
        table.rows[0].cells[2].text = "性别"

        template_path = tmp_path / "test_table.docx"
        doc.save(str(template_path))

        engine = TemplateEngine(sample_resume)
        output_path = tmp_path / "output_table.docx"
        engine.fill_template(template_path, output_path)

    def test_fill_all(self, sample_resume, generated_templates):
        """Test batch filling all templates."""
        engine = TemplateEngine(sample_resume)
        output_dir = generated_templates.parent / "output"
        results = engine.fill_all(generated_templates, output_dir)

        assert len(results) == 3
        for path in results:
            assert path.exists()

    def test_list_keys(self, sample_resume):
        engine = TemplateEngine(sample_resume)
        keys = engine.list_available_keys()
        assert "name" in keys
        assert "phone" in keys
        assert "email" in keys


# ──────────────────── Detector Tests ────────────────────

class TestFieldDetector:
    def test_detect_fields_in_table(self, generated_templates):
        """Test detecting fields in a table-based template."""
        detector = FieldDetector()
        results = detector.detect_fields(generated_templates / "求职登记表.docx")

        assert "name" in results
        assert "gender" in results
        assert "phone" in results
        assert "email" in results

    def test_detect_fields_in_paragraphs(self, generated_templates):
        """Test detecting fields in a paragraph-based template."""
        detector = FieldDetector()
        results = detector.detect_fields(generated_templates / "个人简历.docx")

        # The {{placeholder}} style should also be detected
        assert len(results) > 0

    def test_preview_mapping(self, generated_templates):
        """Test preview mapping."""
        detector = FieldDetector()
        previews = detector.preview_mapping(generated_templates / "高校教师应聘表.docx")

        assert len(previews) > 0
        labels = [p["label"] for p in previews]
        assert any("姓名" in label for label in labels)
        assert any("性别" in label for label in labels)

    def test_auto_fill_table_template(self, sample_resume, generated_templates):
        """Test auto-filling a table template (smart mode)."""
        detector = FieldDetector()
        data = sample_resume.to_flat_dict()
        output_path = generated_templates.parent / "auto_filled.docx"

        detector.auto_fill(
            generated_templates / "求职登记表.docx",
            data,
            output_path,
        )

        assert output_path.exists()

        # Read back and check some fields
        result_doc = Document(str(output_path))
        full_text = " ".join(
            cell.text for table in result_doc.tables
            for row in table.rows
            for cell in row.cells
        )
        assert "测试用户" in full_text or "男" in full_text

    def test_auto_fill_paragraph_template(self, sample_resume, generated_templates):
        """Test auto-filling a paragraph template."""
        detector = FieldDetector()
        data = sample_resume.to_flat_dict()
        output_path = generated_templates.parent / "auto_filled2.docx"

        detector.auto_fill(
            generated_templates / "个人简历.docx",
            data,
            output_path,
        )

        assert output_path.exists()

        result_doc = Document(str(output_path))
        full_text = " ".join(p.text for p in result_doc.paragraphs)
        assert "测试用户" in full_text
        assert "13900000000" in full_text

    def test_auto_fill_all_templates(self, sample_resume, generated_templates):
        """Test auto-filling all templates."""
        detector = FieldDetector()
        data = sample_resume.to_flat_dict()
        output_dir = generated_templates.parent / "batch_output"
        output_dir.mkdir(exist_ok=True)

        for template in sorted(generated_templates.glob("*.docx")):
            out = output_dir / f"filled_{template.name}"
            detector.auto_fill(template, data, out)
            assert out.exists()


# ──────────────────── Template Generator Tests ────────────────────

class TestTemplateGenerator:
    def test_generate_all(self, tmp_path):
        paths = generate_all_templates(tmp_path / "templates")
        assert len(paths) == 3
        for p in paths:
            assert p.exists()
            assert p.suffix == ".docx"

    def test_template_1_has_table(self, tmp_path):
        from auto_resume.template_generator import create_template_1
        path = create_template_1(tmp_path / "t1.docx")
        doc = Document(str(path))
        assert len(doc.tables) > 0

    def test_template_2_has_placeholders(self, tmp_path):
        from auto_resume.template_generator import create_template_2
        path = create_template_2(tmp_path / "t2.docx")
        doc = Document(str(path))
        full_text = " ".join(p.text for p in doc.paragraphs)
        assert "{{name}}" in full_text or "{{" in full_text

    def test_template_3_has_table(self, tmp_path):
        from auto_resume.template_generator import create_template_3
        path = create_template_3(tmp_path / "t3.docx")
        doc = Document(str(path))
        assert len(doc.tables) > 0
